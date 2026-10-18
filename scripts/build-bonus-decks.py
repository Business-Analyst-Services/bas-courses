"""Builds blank + worked-example decks for bonus playbooks B1–B3 in the house style,
reusing the card layout of 04_Requirements. Run: python3 scripts/build-bonus-decks.py"""
import copy
from pptx import Presentation

DIR = "private/files/ba-playbooks/"
BASE = DIR + "04_Requirements_Worked_Example.pptx"

# Each step: (step title, artefact, [ (item, tag, body, source) x5 ])
BONUS = {
 "B1": ("BA Job Search & Interview", "B1_Job_Search", [
  ("Frame the Target Role", "Positioning Brief", [
   ("Target role", "Role", "Senior BA, financial services or payments transformation, Sydney hybrid", "Source: Career Outcome Statement"),
   ("Value I move", "Value", "Time-to-decision and cost-to-serve: cut onboarding 14 → 5 days at Northbridge", "Source: Evidence Inventory"),
   ("Positioning line", "Pitch", "I reframe problems as decisions, then deliver the change that moves the number", "Source: drafted, BA-approved"),
   ("Not targeting", "Filter", "Pure documentation roles; on-site only; contracts under 6 months", "Source: Career Outcome Statement"),
   ("Proof needed", "Gap", "One sanitised business case and one BPMN To-Be to show in interview", "Source: gap flagged by agent"),
  ]),
  ("Map the Hiring Stakeholders", "Target & Contact Map", [
   ("Hiring manager", "Accel.", "Transformation Lead, Acme Bank. Needs a BA who can run Discovery solo", "Warm path: ex-colleague on her team"),
   ("Recruiter", "Neutral", "Agency recruiter screening on keywords: BPMN, payments, stakeholder mgmt", "Action: send tailored resume first"),
   ("Product owner", "Neutral", "Payments PO; will sit on panel and test story-writing", "Action: prepare a Given/When/Then sample"),
   ("Head of Ops", "Blocker?", "Burned by a previous BA who never left the room; sceptical", "Action: lead with an ops outcome story"),
   ("Peer BA", "Accel.", "Existing BA on team; can describe the real problem before interview", "Action: 20-min coffee chat request"),
  ]),
  ("Build the Evidence Portfolio", "Evidence Portfolio + STAR bank", [
   ("Business cases", "Must", "One-page case (sanitised). Outcome: $890k NPV approved", "Contribution: framed options, ran scoring"),
   ("Process mapping", "Must", "As-Is / To-Be onboarding BPMN; 3 waits removed", "Contribution: led workshops, drew model"),
   ("Stakeholder mgmt", "Must", "STAR: turned compliance from blocker to co-designer in 2 weeks", "Contribution: decision register"),
   ("Data & KPIs", "Should", "KPI tree tying drop-off to 3 value leaks", "Contribution: analysis + leak sizing"),
   ("Agile delivery", "Should", "Story backlog with traceability to requirements", "Gap: no cloud evidence — flagged"),
  ]),
  ("Run the Interview & Offer Loop", "Interview Plan + Offer Decision", [
   ("Opening pitch", "Pitch", "'I cut onboarding from 14 to 5 days by reframing the decisions — here's how I'd approach yours'", "Use: first 60 seconds"),
   ("Likely question", "Prep", "'Tell me about a difficult stakeholder' → compliance STAR story", "Evidence: STAR bank #3"),
   ("My questions", "Prep", "Which decision is blocking delivery today? Who owns it?", "Purpose: show Discovery thinking"),
   ("Offer criteria", "Decide", "Scope to own Discovery; hybrid; ≥ target salary band", "Walk away if 2 of 3 missing"),
   ("Follow-up", "Action", "Thank-you note within 24h with one-page approach sketch", "Owner: me · Due: next day"),
  ]),
 ]),
 "B2": ("Application Tailoring & Fit Assessment", "B2_Application_Tailoring", [
  ("Deconstruct the Job Ad", "Job Ad Breakdown", [
   ("Mandate", "Role", "Own end-to-end Discovery for a payments transformation", "Source: ad paragraph 1"),
   ("Must-haves", "Must", "5+ yrs BA, BPMN, stakeholder management", "Source: requirements list"),
   ("Nice-to-haves", "Should", "Payments domain, Jira/Confluence, agile ceremonies", "Source: 'desirable' list"),
   ("Keywords", "ATS", "elicitation, user stories, BPMN, business case, payments", "Use: mirror in resume"),
   ("Disqualifier", "Check", "Must hold current work rights — confirm before applying", "Source: ad footer"),
  ]),
  ("Compare Against Resume & Cover Letter", "Fit Comparison Matrix", [
   ("BPMN mapping", "Strong", "'Redesigned onboarding As-Is / To-Be' (Northbridge role)", "Keyword match: Y"),
   ("Stakeholder mgmt", "Strong", "Decision register; compliance co-design story", "Keyword match: Y"),
   ("Payments domain", "Partial", "Onboarding, not payments — adjacent fintech", "Keyword match: N — reword"),
   ("Business cases", "Strong", "$890k NPV one-page case approved", "Keyword match: Y"),
   ("Agile tooling", "Partial", "Jira used; Confluence not mentioned", "Keyword match: N — add"),
  ]),
  ("Score Suitability & Identify Gaps", "Suitability Assessment", [
   ("Fit score", "78/100", "Strong on core BA skills; lighter on payments domain", "Verdict: apply with tailoring"),
   ("Gap 1", "Close", "Low payments keyword density — reword two roles", "Closeable now"),
   ("Gap 2", "Close", "Confluence not named — add to tools line", "Closeable now"),
   ("Gap 3", "Explain", "No direct payments role — address in cover letter", "Frame as adjacent fintech"),
   ("Decision", "Apply", "Apply this week; tailored resume + letter", "Owner: me"),
  ]),
  ("Produce Tailored Application Material", "Tailored Resume & Cover Letter", [
   ("Summary", "Before", "'Experienced BA.'", "Weak: no outcome, no domain"),
   ("Summary", "After", "'Senior BA who cut fintech onboarding from 14 to 5 days — seeking to own Discovery on your payments transformation.'", "Mirrors mandate"),
   ("Role bullet", "After", "Mapped As-Is / To-Be onboarding in BPMN; removed 3 waits", "Adds keyword: BPMN"),
   ("Cover letter hook", "After", "Opens with their Discovery mandate, then one proof story", "Under 250 words"),
   ("Review check", "BA owns", "No invented claims; every line traceable to evidence", "I review before sending"),
  ]),
 ]),
 "B3": ("AI Opportunity Match & Pipeline", "B3_Opportunity_Pipeline", [
  ("Build the Profile & Match Model", "Match Model", [
   ("Weights", "Model", "Skills 35 · Domain 20 · Level 15 · Location 15 · Salary 15", "Total 100"),
   ("Hard filter", "Reject", "On-site only roles", "Applied before scoring"),
   ("Band: apply now", "80+", "Apply within 48 hours", "Effort: low"),
   ("Band: tailor", "65–79", "Tailor using Playbook B2 first", "Effort: medium"),
   ("Band: watch", "50–64", "Save; re-check in two weeks", "Below 50: discard"),
  ]),
  ("Connect the App & Pull Roles", "Normalised Role List", [
   ("Acme Bank", "Kept", "Senior BA · Sydney / Hybrid · $140–160k", "BPMN, payments, stakeholder mgmt"),
   ("Harbour Insurance", "Kept", "BA · Hybrid · $130–145k", "claims, process, Jira"),
   ("Metro Health", "Kept", "Lead BA · Hybrid · $150–170k", "requirements, change, data"),
   ("Coastal Retail", "Dropped", "BA · On-site only", "Hard filter"),
   ("Southern Freight", "Kept", "Process Analyst · Hybrid · $120–135k", "BPMN, logistics"),
  ]),
  ("Score & Rank the Opportunities", "Scored Opportunity Board", [
   ("Acme Bank", "86", "Skills 32/35 · Domain 18/20 · Salary 12/15", "Band: apply now · Effort: low"),
   ("Metro Health", "79", "Skills 30/35 · Domain 12/20 · Salary 15/15", "Band: tailor · Gap: health domain"),
   ("Harbour Insurance", "72", "Skills 29/35 · Domain 14/20 · Salary 10/15", "Band: tailor"),
   ("Southern Freight", "61", "Skills 27/35 · Level 8/15 · Salary 8/15", "Band: watch"),
   ("Top gap overall", "Note", "Light on cloud platforms across 3 of 4 roles", "Action: short course"),
  ]),
  ("Prioritise & Trigger Applications", "Application Action Plan", [
   ("This week", "Apply", "Acme Bank — tailored resume, warm intro via ex-colleague", "Due: Wednesday"),
   ("This week", "Tailor", "Metro Health, Harbour Insurance — run Playbook B2", "Due: Friday"),
   ("Watch", "Hold", "Southern Freight — re-score if level changes", "Re-check: 2 weeks"),
   ("Model tune", "Adjust", "No callbacks below $150k after 2 weeks → raise salary weight", "Review: fortnightly"),
   ("Guardrail", "BA owns", "Agent drafts; I approve every application before it is sent", "Never auto-submit"),
  ]),
 ]),
}

def dup_slide(prs, src):
    s = prs.slides.add_slide(src.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    for el in src.shapes._spTree:
        if el.tag.endswith(("}sp", "}pic", "}grpSp", "}cxnSp")):
            s.shapes._spTree.append(copy.deepcopy(el))
    for rel in src.part.rels.values():
        if "image" in rel.reltype:
            s.part.rels.get_or_add(rel.reltype, rel._target)
    return s

def set_text(shape, text):
    tf = shape.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.text = text
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

def build(code, blank):
    title, fname, steps = BONUS[code]
    prs = Presentation(BASE)
    src = prs.slides[0]
    # drop slides 2-4, then duplicate slide 1 three times
    ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]:
        prs.part.drop_rel(sid.rId); ids.remove(sid)
    for _ in range(3):
        dup_slide(prs, src)
    for slide, (stitle, artefact, cards) in zip(prs.slides, steps):
        tb = [sh for sh in slide.shapes if sh.has_text_frame]
        by = {sh.name: sh for sh in tb}
        n = list(prs.slides).index(slide) + 1
        set_text(by["TextBox 3"], f"BONUS PLAYBOOK {code}  ·  {title}")
        set_text(by["TextBox 4"], f"Step {n} — {stitle}")
        set_text(by["TextBox 5"], f"ARTEFACT  ·  {artefact}" + (" — blank template" if blank else ""))
        set_text(by["TextBox 6"], "YOUR PROJECT" if blank else "WORKED EXAMPLE")
        set_text(by["TextBox 47"], "Bonus playbooks · " + ("YOUR SEARCH" if blank else "Fictional candidate"))
        # remove marquee badge
        for nm in ("Rounded Rectangle 48", "TextBox 49"):
            for sh in slide.shapes:
                if sh.name == nm:
                    sh._element.getparent().remove(sh._element)
        groups = [("TextBox 11", "TextBox 13", "TextBox 14", "TextBox 15"),
                  ("TextBox 18", "TextBox 20", "TextBox 21", "TextBox 22"),
                  ("TextBox 25", "TextBox 27", "TextBox 28", "TextBox 29"),
                  ("TextBox 32", "TextBox 34", "TextBox 35", "TextBox 36"),
                  ("TextBox 39", "TextBox 41", "TextBox 42", "TextBox 43")]
        for g, (item, tag, body, srcl) in zip(groups, cards):
            vals = (item, "<tag>", "…", "…") if blank else (item, tag, body, srcl)
            for nm, v in zip(g, vals):
                set_text(by[nm], v)
        # widen tag pills to fit
        for sh in slide.shapes:
            if sh.name in ("Rounded Rectangle 12", "Rounded Rectangle 19", "Rounded Rectangle 26",
                           "Rounded Rectangle 33", "Rounded Rectangle 40",
                           "TextBox 13", "TextBox 20", "TextBox 27", "TextBox 34", "TextBox 41"):
                sh.width = 12700 * 64
    out = f"{DIR}{fname}_{'Blank_Template' if blank else 'Worked_Example'}.pptx"
    prs.save(out)
    print("wrote", out)

for code in BONUS:
    build(code, True)
    build(code, False)
